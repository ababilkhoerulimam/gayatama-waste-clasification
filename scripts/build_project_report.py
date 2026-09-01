from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Laporan_Gagasan_AI_Klasifikasi_Sampah.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
GRAY = "666666"
LIGHT_FILL = "F4F6F9"
MID_FILL = "E8EEF5"
WHITE = "FFFFFF"
BLACK = "000000"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "19")
    r_pr.append(color)
    r_pr.append(underline)
    r_pr.append(size)
    r_pr.append(size_cs)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    paragraph.add_run("Halaman ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, display, end])


def add_toc_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "Daftar isi diperbarui otomatis saat dokumen dibuka di Microsoft Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, fallback, end])


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead, rest = text.split(":", 1)
        r = p.add_run(lead + ":")
        r.bold = True
        p.add_run(rest)
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbers(doc, items):
    num_id = new_numbering_id(doc)
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_paragraph_num_id(p, num_id)
        p.add_run(item)


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lead = p.add_run(label + " ")
    lead.bold = True
    lead.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        set_cell_shading(header.cells[idx], LIGHT_FILL)
        p = header.cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, size=font_size, color=NAVY, bold=True)
    for row_values in rows:
        row = table.add_row()
        for idx, text in enumerate(row_values):
            p = row.cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            if idx == 0 and len(headers) == 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(text))
            set_run_font(run, size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def new_numbering_id(doc):
    numbering = doc.part.numbering_part.element
    style_num_id = doc.styles["List Number"]._element.pPr.numPr.numId.val
    source_num = next(
        el for el in numbering.findall(qn("w:num"))
        if el.get(qn("w:numId")) == str(style_num_id)
    )
    abstract_id = source_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    used_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    new_id = max(used_ids) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_id


def set_paragraph_num_id(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")
    num = num_pr.find(qn("w:numId"))
    if num is None:
        num = OxmlElement("w:numId")
        num_pr.append(num)
    num.set(qn("w:val"), str(num_id))


def add_reference(doc, num_id, citation, url):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_num_id(p, num_id)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    citation_run = p.add_run(citation + " ")
    set_run_font(citation_run, size=9.5)
    add_hyperlink(p, "Tautan sumber", url)


def add_table_caption(doc, text):
    p = doc.add_paragraph(text, style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    return p


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    doc.core_properties.title = "Gagasan Pemanfaatan AI untuk Klasifikasi dan Pengarahan Sampah"
    doc.core_properties.subject = "Laporan gagasan proyek klasifikasi sampah berbasis computer vision"
    doc.core_properties.author = "Tim Proyek"
    doc.core_properties.keywords = "AI, klasifikasi sampah, PyTorch, human-in-the-loop, waste management"

    # Running header/footer for body pages.
    hp = section.header.paragraphs[0]
    hp.text = "LAPORAN GAGASAN PROYEK  |  AI UNTUK KLASIFIKASI SAMPAH"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    for run in hp.runs:
        set_run_font(run, size=8.5, color=GRAY, bold=True)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    add_page_field(fp)
    for run in fp.runs:
        set_run_font(run, size=8.5, color=GRAY)

    # Editorial cover.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(76)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    set_run_font(kicker.add_run("LAPORAN GAGASAN PROYEK"), size=11, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_run_font(title.add_run("PEMANFAATAN AI UNTUK\nKLASIFIKASI DAN PENGARAHAN SAMPAH"), size=25, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(46)
    set_run_font(subtitle.add_run("Proof of Concept Berbasis PyTorch dengan Human-in-the-Loop"), size=14, color=DARK_BLUE, italic=True)

    meta_lines = [
        "Disusun oleh: [Nama Anggota Kelompok]",
        "Mata kuliah: [Nama Mata Kuliah]",
        "Dosen pengampu: [Nama Dosen]",
        "Institusi: [Nama Institusi]",
    ]
    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run(line), size=11, color=GRAY)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(34)
    set_run_font(date.add_run("2026"), size=12, color=NAVY, bold=True)

    doc.add_page_break()

    add_heading(doc, "Ringkasan Eksekutif", 1)
    add_body(doc, "Laporan ini mengusulkan pemanfaatan kecerdasan buatan berbasis computer vision untuk membantu proses pemilahan sampah. Foto atau frame kamera digunakan sebagai masukan, kemudian model PyTorch mengidentifikasi satu objek sampah dominan ke dalam sepuluh kelas material. Hasil tersebut diproses oleh lapisan aturan untuk memberikan rekomendasi awal berupa kandidat bahan furnitur, kandidat pemulihan energi, atau penanganan lain/residu.")
    add_body(doc, "Sistem dirancang sebagai decision support, bukan pengganti manusia. Prediksi ber-confidence rendah, ambigu, atau berisiko tinggi diarahkan ke status NEEDS_HUMAN_REVIEW. Petugas dapat menyetujui atau memperbaiki hasil model. Pendekatan ini menggambarkan bagaimana AI berpotensi membuat alur pemilahan lebih terstruktur sambil mempertahankan pengawasan manusia.")
    add_callout(doc, "Posisi penelitian:", "proof of concept berbasis dataset dan evaluasi statistik. Penelitian belum menguji dampak operasional langsung pada fasilitas pengolahan sampah.")

    add_heading(doc, "Daftar Isi", 1)
    toc = doc.add_paragraph()
    add_toc_field(toc)
    doc.add_page_break()

    add_heading(doc, "1. Pendahuluan", 1)
    add_heading(doc, "1.1 Latar Belakang", 2)
    add_body(doc, "Pemilahan menentukan apakah suatu material dapat dimanfaatkan kembali, dipulihkan energinya, memerlukan pengolahan lain, atau menjadi residu. Dalam proses manual, petugas harus mengidentifikasi material dan menentukan jalurnya satu per satu. Aktivitas berulang ini berpotensi menghasilkan beban kerja tinggi, keputusan yang kurang konsisten, serta risiko salah penanganan pada benda ambigu atau berbahaya.")
    add_body(doc, "AI berbasis computer vision dapat memberi identifikasi dan rekomendasi awal. Namun, kamera biasa tidak dapat memastikan seluruh karakteristik teknis material, seperti kadar air, nilai kalor, klorin, logam berat, atau kontaminasi kimia. Karena itu, hasil sistem harus dibaca sebagai rekomendasi awal yang dilengkapi human-in-the-loop.")

    add_heading(doc, "1.2 Pernyataan Masalah", 2)
    add_body(doc, "Bagaimana model deep learning berbasis citra dapat digunakan sebagai proof of concept untuk mengklasifikasikan material sampah, menampilkan confidence score, memberi rekomendasi jalur penanganan, dan menyerahkan kasus yang tidak pasti atau berisiko kepada manusia?")

    add_heading(doc, "1.3 Tujuan", 2)
    add_bullets(doc, [
        "Mengembangkan proof of concept klasifikasi citra sampah menggunakan PyTorch.",
        "Mengelompokkan sampah ke dalam sepuluh kelas material yang dapat dipetakan ke rekomendasi penanganan.",
        "Membandingkan model menggunakan metrik klasifikasi dan kebutuhan komputasi.",
        "Menguji penggunaan confidence threshold sebagai dasar rujukan kepada manusia.",
        "Menyusun gagasan implementasi AI yang realistis, aman, dan mudah dijelaskan kepada masyarakat.",
    ])

    add_heading(doc, "1.4 Pertanyaan Penelitian", 2)
    add_numbers(doc, [
        "Seberapa baik model membedakan sepuluh kelas material pada dataset penelitian?",
        "Model mana yang memberi keseimbangan terbaik antara performa dan kebutuhan komputasi?",
        "Bagaimana confidence threshold memengaruhi automatic coverage dan referral rate?",
        "Kesalahan klasifikasi apa yang paling sering muncul dan apa dampaknya terhadap rekomendasi?",
        "Bagaimana human-in-the-loop dapat mengurangi risiko keputusan otomatis?",
    ])

    add_heading(doc, "2. Ruang Lingkup dan Batas Penelitian", 1)
    add_heading(doc, "2.1 Termasuk dalam Ruang Lingkup", 2)
    add_bullets(doc, [
        "Dataset citra publik dan/atau foto tambahan yang dikumpulkan tim.",
        "Klasifikasi satu objek sampah dominan dalam satu gambar atau frame kamera.",
        "Pelatihan baseline dan model transfer learning menggunakan PyTorch.",
        "Simulasi rekomendasi jalur penanganan dan human-in-the-loop.",
        "Evaluasi statistik, demonstrasi foto/kamera, paper, serta video edukasi.",
    ])
    add_heading(doc, "2.2 Di Luar Ruang Lingkup", 2)
    add_bullets(doc, [
        "Pemasangan sistem pada fasilitas nyata atau integrasi dengan conveyor dan robot.",
        "Pengujian laboratorium material dan pengukuran langsung efisiensi operasional.",
        "Keputusan otomatis final tanpa verifikasi manusia.",
        "Deteksi banyak objek bertumpuk dalam satu frame pada versi awal.",
        "Klaim bahwa sistem telah terbukti meningkatkan efisiensi fasilitas nyata.",
    ])

    add_heading(doc, "3. Gagasan Sistem", 1)
    add_body(doc, "Sistem konseptual terdiri dari tiga lapisan: model visual, pemeriksaan kondisi/risiko, dan aturan fasilitas. Pemisahan ini penting karena confidence model hanya menunjukkan tingkat keyakinan prediksi kelas; confidence tidak menentukan kelayakan teknis suatu material untuk furnitur atau energi.")
    add_numbers(doc, [
        "Model visual menghasilkan kelas material dan confidence score.",
        "Lapisan pemeriksaan menilai threshold, ambiguitas, kualitas gambar, dan aturan risiko.",
        "Lapisan rekomendasi memetakan kelas serta kondisi yang tersedia ke jalur penanganan awal.",
        "Petugas memverifikasi kasus yang dirujuk dan dapat memperbaiki hasil.",
    ])
    add_callout(doc, "Alur ringkas:", "foto/kamera -> klasifikasi material -> pemeriksaan confidence dan risiko -> rekomendasi awal atau pemeriksaan manusia.")

    add_heading(doc, "4. Taksonomi Sepuluh Kelas", 1)
    add_body(doc, "Taksonomi mengadaptasi kelompok komposisi sampah SIPSN. Plastik dipisahkan menjadi plastik keras dan fleksibel, sedangkan baterai/elektronik dipisahkan untuk mendukung keselamatan.")
    class_rows = [
        (1, "organic", "Sampah organik", "Sisa makanan, kulit buah"),
        (2, "wood_vegetation", "Kayu dan vegetasi", "Kayu, ranting, daun"),
        (3, "paper_cardboard", "Kertas dan kardus", "Kertas, koran, kardus"),
        (4, "rigid_plastic", "Plastik keras", "Botol, jeriken, wadah keras"),
        (5, "flexible_plastic", "Plastik fleksibel", "Kresek, sachet, plastik film"),
        (6, "textile_rubber_leather", "Tekstil, karet, kulit", "Kain, sepatu, potongan karet"),
        (7, "metal", "Logam", "Kaleng, aluminium, besi"),
        (8, "glass_ceramic", "Kaca dan keramik", "Botol kaca, pecahan, keramik"),
        (9, "battery_electronic", "Baterai dan elektronik", "Baterai, kabel, komponen"),
        (10, "mixed_residual", "Campuran atau residu", "Popok, pembalut, benda multimaterial"),
    ]
    add_table_caption(doc, "Tabel 1. Taksonomi kelas material yang diusulkan.")
    add_table(doc, ["No.", "Label model", "Nama umum", "Contoh"], class_rows, [500, 2450, 2500, 3910], 8.5)

    add_heading(doc, "5. Dasar Tiga Jalur Penanganan", 1)
    add_body(doc, "Tiga jalur merupakan rekomendasi pengolahan, bukan kelas visual. Penentuannya mengikuti waste hierarchy: pemanfaatan material didahulukan, kemudian pemulihan energi apabila material tidak layak dimanfaatkan, dan penanganan lain atau residu menjadi pilihan berikutnya. Jalur final tetap mengikuti spesifikasi fasilitas dan peraturan yang berlaku.")

    add_heading(doc, "5.1 Kriteria Teknis", 2)
    metric_rows = [
        ("Jenis material", "Polimer yang diterima", "Material mudah terbakar", "Material tidak kompatibel"),
        ("Kebersihan", "Kontaminasi rendah", "Bebas bahan berbahaya", "Kotor/berisiko"),
        ("Homogenitas", "Tinggi/kompatibel", "Campuran tertentu", "Tidak diketahui/kompleks"),
        ("Kadar air", "Sebaiknya rendah", "Kritis; semakin rendah semakin baik", "Terlalu basah"),
        ("Nilai kalor", "Tidak menjadi dasar", "Parameter utama", "Tidak memenuhi spesifikasi"),
        ("Kadar abu", "Tidak menjadi dasar", "Perlu dikendalikan", "Terlalu tinggi"),
        ("Klorin/halogen", "Sesuai proses", "Perlu dibatasi", "Melebihi spesifikasi"),
        ("Logam berat", "Tidak membahayakan produk", "Perlu dibatasi", "Penanganan khusus"),
        ("Penerimaan fasilitas", "Wajib diterima", "Wajib diterima", "Ditolak kedua jalur"),
    ]
    add_table_caption(doc, "Tabel 2. Kriteria konseptual pemilihan jalur penanganan.")
    add_table(doc, ["Kriteria", "Furnitur", "Pemulihan energi", "Penanganan lain"], metric_rows, [1500, 2450, 2910, 2500], 8.2)

    add_heading(doc, "5.2 Kandidat Bahan Furnitur", 2)
    add_body(doc, "Kandidat furnitur memerlukan material yang kompatibel dengan proses, relatif bersih, homogen, tidak berbahaya, dan dapat diproses menjadi produk dengan sifat mekanis yang memadai. ASTM D7568 mengizinkan sistem polyethylene daur ulang untuk structural plastic lumber, tetapi produk akhirnya tetap harus memenuhi persyaratan teknis. Model gambar hanya dapat memberi kandidat awal; jenis polimer dan kualitas produk memerlukan verifikasi tambahan.")

    add_heading(doc, "5.3 Kandidat Pemulihan Energi", 2)
    add_body(doc, "Kelayakan bahan bakar sampah ditentukan oleh karakteristik seperti nilai kalor, kadar air, kadar abu, klorin, sulfur, merkuri, ukuran partikel, dan status non-B3. ISO 21640 menegaskan bahwa sampah kota mentah bukan langsung solid recovered fuel; sampah harus melalui pengolahan. Studi RDF Plant Cilacap juga mengevaluasi kadar air, kadar abu, dan nilai kalor serta menunjukkan peran biodrying dalam meningkatkan kualitas bahan bakar.")

    add_heading(doc, "5.4 Penanganan Lain atau Residu", 2)
    add_body(doc, "Kategori ini mencakup material yang masih memiliki jalur lain, seperti daur ulang logam/kaca dan pengolahan organik, serta material yang membutuhkan penanganan khusus. Karena itu, OTHER_HANDLING_OR_RESIDUAL tidak identik dengan dibuang. Baterai, elektronik, material berbahaya, dan objek tidak dikenal harus dirujuk kepada manusia.")

    add_heading(doc, "5.5 Matriks Pemetaan Awal", 2)
    route_rows = [
        ("organic", "Penanganan lain", "Pengomposan/pengolahan biologis"),
        ("wood_vegetation", "Energi", "Jika cukup kering dan diterima fasilitas"),
        ("paper_cardboard", "Energi/penanganan lain", "Daur ulang jika bersih; energi jika tidak layak didaur ulang dan kering"),
        ("rigid_plastic", "Furnitur", "Tergantung polimer, kebersihan, dan fasilitas"),
        ("flexible_plastic", "Furnitur/energi", "Tergantung teknologi dan spesifikasi fasilitas"),
        ("textile_rubber_leather", "Energi", "Jika jenis material dan fasilitas mengizinkan"),
        ("metal", "Penanganan lain", "Daur ulang logam; bukan bahan bakar"),
        ("glass_ceramic", "Penanganan lain", "Daur ulang/residu; bukan bahan bakar"),
        ("battery_electronic", "Penanganan khusus", "Wajib diperiksa manusia"),
        ("mixed_residual", "Penanganan lain/residu", "Wajib diperiksa sebelum keputusan akhir"),
    ]
    add_table_caption(doc, "Tabel 3. Pemetaan awal kelas material ke rekomendasi penanganan.")
    add_table(doc, ["Kelas", "Rekomendasi awal", "Catatan"], route_rows, [2500, 2300, 4560], 8.4)

    add_heading(doc, "6. Human-in-the-Loop", 1)
    add_body(doc, "Human-in-the-loop memastikan AI berfungsi sebagai pendukung keputusan. NEEDS_HUMAN_REVIEW bukan kelas sampah ke-11, melainkan status yang diberikan setelah prediksi model diperiksa oleh lapisan keputusan.")
    hitl_rows = [
        ("AUTO_ACCEPTED", "Confidence melewati threshold dan tidak terkena aturan risiko"),
        ("NEEDS_HUMAN_REVIEW", "Confidence rendah, ambigu, berisiko, atau di luar kelas yang dikenal"),
        ("HUMAN_CONFIRMED", "Petugas menyetujui prediksi dan rekomendasi"),
        ("HUMAN_CORRECTED", "Petugas mengganti kelas atau rekomendasi"),
    ]
    add_table_caption(doc, "Tabel 4. Status keputusan dalam alur human-in-the-loop.")
    add_table(doc, ["Status", "Makna"], hitl_rows, [3000, 6360], 9)

    add_heading(doc, "6.1 Aturan Rujukan", 2)
    add_bullets(doc, [
        "Confidence prediksi tertinggi berada di bawah threshold tau.",
        "Selisih confidence antara dua kelas teratas terlalu kecil.",
        "Model memprediksi battery_electronic atau mixed_residual.",
        "Gambar buruk, objek tertutup, atau lebih dari satu objek dominan terlihat.",
        "Objek berada di luar kelompok yang dikenal model.",
        "Aturan keselamatan fasilitas mewajibkan pemeriksaan manusia.",
    ])
    add_body(doc, "Nilai threshold dipilih menggunakan validation set dengan membandingkan automatic coverage, referral rate, dan selective accuracy. Softmax diperlakukan sebagai confidence score, bukan jaminan probabilitas bahwa prediksi benar.")

    add_heading(doc, "6.2 Umpan Balik Manusia", 2)
    add_body(doc, "Koreksi petugas dapat disimpan sebagai calon data berlabel. Sebelum digunakan untuk pelatihan ulang, data harus melalui pemeriksaan kualitas label. Langkah ini menghindari masuknya koreksi yang keliru atau tidak konsisten ke dalam dataset.")

    add_heading(doc, "7. Rencana Metodologi", 1)
    add_heading(doc, "7.1 Pemilihan Dataset", 2)
    add_bullets(doc, [
        "Lisensi dan sumber dataset harus jelas.",
        "Label dapat dipetakan secara konsisten ke taksonomi proyek.",
        "Gambar memiliki variasi cahaya, sudut, kondisi objek, dan latar.",
        "Dataset tidak hanya berisi objek pada latar studio yang terlalu bersih.",
        "Training, validation, dan test dapat dipisahkan tanpa data leakage.",
        "Perbedaan sumber dataset diperiksa agar model tidak belajar mengenali latar atau sumber data.",
    ])

    add_heading(doc, "7.2 Model dan Eksperimen", 2)
    add_body(doc, "Eksperimen mencakup satu CNN sederhana sebagai baseline dan satu atau beberapa model pretrained dengan transfer learning. Pemilihan arsitektur final dilakukan setelah dataset tersedia. Kontribusi proyek tidak harus berupa arsitektur baru; kontribusi utamanya dapat berupa integrasi klasifikasi, rekomendasi berbasis aturan, dan human-in-the-loop.")

    add_heading(doc, "7.3 Pembagian Data", 2)
    add_body(doc, "Dataset dibagi menjadi training set, validation set, dan test set. Pembagian dibuat stratified jika memungkinkan. Gambar dari objek atau rangkaian pengambilan yang sama tidak boleh tersebar ke training dan test set karena dapat menyebabkan data leakage.")

    add_heading(doc, "8. Rencana Evaluasi", 1)
    eval_rows = [
        ("Klasifikasi", "Accuracy; precision, recall, dan F1-score per kelas; macro average; confusion matrix"),
        ("Efisiensi model", "Waktu inferensi per gambar dan ukuran model"),
        ("Human-in-the-loop", "Automatic coverage, referral rate, selective accuracy, error referral rate"),
        ("Analisis", "Perbandingan threshold dan analisis jenis kesalahan yang berisiko"),
    ]
    add_table_caption(doc, "Tabel 5. Kelompok metrik evaluasi penelitian.")
    add_table(doc, ["Aspek", "Metrik"], eval_rows, [2500, 6860], 9)
    add_body(doc, "Macro average dan metrik per kelas diprioritaskan bersama accuracy agar performa pada kelas kecil tidak tertutup oleh kelas dominan. Threshold terbaik bukan semata-mata threshold dengan referral rate paling rendah, melainkan yang mempertahankan selective accuracy tinggi pada tingkat pemeriksaan manusia yang masih masuk akal.")

    add_heading(doc, "9. Skenario Demonstrasi", 1)
    add_numbers(doc, [
        "Pengguna mengunggah foto atau mengarahkan satu objek ke kamera.",
        "Sistem menampilkan kelas material dan confidence score.",
        "Lapisan keputusan memeriksa threshold dan aturan risiko.",
        "Sistem menampilkan rekomendasi awal atau status NEEDS_HUMAN_REVIEW.",
        "Petugas menyetujui atau memperbaiki kelas dan rekomendasi.",
        "Hasil keputusan ditampilkan dan dapat dicatat sebagai riwayat simulasi.",
    ])
    add_callout(doc, "Batas demonstrasi:", "demonstrasi menunjukkan alur pendukung keputusan dan bukan replika penuh fasilitas industri.")

    add_heading(doc, "10. Luaran Proyek", 1)
    add_heading(doc, "10.1 Paper", 2)
    add_body(doc, "Paper memuat latar belakang, studi terkait, taksonomi kelas, rancangan human-in-the-loop, dataset, preprocessing, model, eksperimen, hasil statistik, analisis kesalahan, simulasi rekomendasi, keterbatasan, dan peluang pengembangan.")
    add_heading(doc, "10.2 Proof of Concept", 2)
    add_bullets(doc, [
        "Model PyTorch terlatih dan kode pelatihan/evaluasi.",
        "Inferensi melalui foto dan, jika memungkinkan, kamera.",
        "Tampilan kelas, confidence, rekomendasi, dan status pemeriksaan manusia.",
    ])
    add_heading(doc, "10.3 Video Edukasi", 2)
    add_body(doc, "Video menjelaskan masalah pemilahan manual, cara sederhana AI mengenali gambar, sepuluh kelas sampah, tiga kelompok rekomendasi, peran manusia, manfaat potensial, dan keterbatasan sistem dengan bahasa umum.")
    add_callout(doc, "Pesan utama:", "AI dapat membantu mengenali dan mengarahkan sampah secara lebih terstruktur, tetapi keputusan penting tetap membutuhkan manusia dan aturan fasilitas yang sesuai.")

    add_heading(doc, "11. Tahapan Pekerjaan", 1)
    add_numbers(doc, [
        "Memfinalkan definisi dan contoh setiap kelas.",
        "Mencari serta mengaudit dataset publik.",
        "Menetapkan aturan pemetaan kelas ke jalur rekomendasi.",
        "Menyiapkan preprocessing dan pembagian data.",
        "Melatih baseline dan model transfer learning.",
        "Mengevaluasi performa serta pola kesalahan.",
        "Memilih dan menguji confidence threshold.",
        "Membuat simulasi human-in-the-loop.",
        "Menyusun paper berdasarkan hasil eksperimen.",
        "Membuat video edukasi berdasarkan gagasan dan hasil penelitian.",
    ])

    add_heading(doc, "12. Batasan, Risiko, dan Etika", 1)
    add_bullets(doc, [
        "Jenis polimer tidak selalu dapat dikenali dari tampilan visual.",
        "Kadar air, nilai kalor, klorin, logam berat, dan kontaminasi kimia tidak dapat dipastikan melalui kamera biasa.",
        "Confidence tinggi tidak menjamin prediksi benar.",
        "Dataset publik mungkin tidak merepresentasikan kondisi sampah Indonesia atau fasilitas nyata.",
        "Model dapat mempelajari shortcut dari latar belakang atau sumber dataset.",
        "Pemetaan jalur berbeda antarwilayah, teknologi, dan fasilitas.",
        "Sistem tidak boleh mendorong masyarakat membakar sampah sendiri.",
        "Material berbahaya dan objek tidak dikenal wajib dirujuk kepada manusia.",
        "Klaim efisiensi aktual memerlukan pengujian lapangan terpisah.",
    ])

    add_heading(doc, "13. Kesimpulan", 1)
    add_body(doc, "Gagasan ini menempatkan AI sebagai alat bantu untuk mengidentifikasi material sampah, menampilkan confidence, dan memberi rekomendasi awal jalur penanganan. Sepuluh kelas material dipisahkan dari tiga rekomendasi agar sistem tidak menyamakan jenis benda dengan kelayakan teknis pengolahannya. Human-in-the-loop menjaga agar kasus ragu, berisiko, dan tidak dikenal tetap diperiksa manusia.")
    add_body(doc, "Dalam ruang lingkup saat ini, keluaran utama adalah proof of concept dan bukti statistik pada dataset, bukan sistem industri. Jika hasil awal baik, penelitian dapat dilanjutkan melalui kalibrasi confidence, integrasi sensor tambahan, object detection multiobjek, dan validasi bersama fasilitas pengolahan sampah.")

    add_heading(doc, "14. Pertanyaan Terbuka", 1)
    add_bullets(doc, [
        "Material apa yang benar-benar diterima dan ditolak oleh fasilitas yang menjadi inspirasi proyek?",
        "Dataset publik mana yang paling sesuai dengan sepuluh kelas dan konteks gambar dunia nyata?",
        "Apakah kelas tekstil, karet, dan kulit perlu dipisahkan setelah dataset ditemukan?",
        "Threshold berapa yang memberi keseimbangan terbaik antara automatic coverage dan selective accuracy?",
        "Apakah koreksi manusia hanya disimulasikan atau disimpan untuk pengembangan berikutnya?",
    ])

    add_heading(doc, "Daftar Pustaka", 1)
    refs = [
        ("Kementerian Lingkungan Hidup dan Kehutanan. Sistem Informasi Pengelolaan Sampah Nasional: Komposisi Sampah.", "https://sipsn.menlhk.go.id/sipsn/public/data/komposisi"),
        ("Pemerintah Republik Indonesia. Peraturan Presiden Nomor 109 Tahun 2025 tentang Penanganan Sampah Perkotaan Melalui Pengolahan Sampah Menjadi Energi Terbarukan Berbasis Teknologi Ramah Lingkungan.", "https://peraturan.bpk.go.id/Details/334718"),
        ("Pemerintah Republik Indonesia. Undang-Undang Nomor 18 Tahun 2008 tentang Pengelolaan Sampah.", "https://peraturan.bpk.go.id/Home/Download/28462/UU%20Nomor%2018%20Tahun%202008.pdf"),
        ("International Organization for Standardization. ISO 21640:2021 Solid recovered fuels - Specifications and classes.", "https://www.iso.org/standard/71309.html"),
        ("ASTM International. ASTM D7568: Standard Specification for Polyethylene-Based Structural-Grade Plastic Lumber for Outdoor Applications.", "https://store.astm.org/d7568-17.html"),
        ("European Commission. Refuse Derived Fuel, Current Practice and Perspectives.", "https://ec.europa.eu/environment/pdf/waste/studies/rdf.pdf"),
        ("Maulidayanti, E. M., dkk. Evaluasi Produksi Refuse-Derived Fuel dari Sampah Perkotaan: Studi Kasus RDF Plant Kabupaten Cilacap. Jurnal Teknologi Lingkungan, BRIN.", "https://ejournal.brin.go.id/JTL/article/view/1008"),
        ("United States Environmental Protection Agency. Decision Maker's Guide to Recycling Plastics.", "https://nepis.epa.gov/Exe/ZyPURL.cgi?Dockey=50000PQU.TXT"),
        ("UCI Machine Learning Repository. RealWaste Data Set.", "https://archive.ics.uci.edu/dataset/908/realwaste"),
    ]
    reference_num_id = new_numbering_id(doc)
    for citation, url in refs:
        add_reference(doc, reference_num_id, citation, url)

    # Update fields when opened in Word.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
